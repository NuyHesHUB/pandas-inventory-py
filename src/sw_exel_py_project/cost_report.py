from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import datetime, timedelta
from difflib import SequenceMatcher
import math
import re
from pathlib import Path
import unicodedata

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import get_column_letter, range_boundaries
import pandas as pd

from .exel_reader import fifo_for_product, load_inventory_df


_CLEAR_CELL = object()
_KEEP_CELL = object()


@dataclass
class TemplateRow:
    excel_row: int
    item: str
    unit_price_raw: object
    unit_price_num: float | None
    exchange_rate: float | None
    duty_rate: float | None
    duty_factor: float | None
    quantity: float
    note: object
    note_date: pd.Timestamp | None


def _canon(s: object) -> str:
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s))
    return re.sub(r"\s+|[()\[\]{}\/\\\-_]", "", s).lower()


def _to_float(v: object) -> float | None:
    try:
        if v is None:
            return None
        if isinstance(v, float) and math.isnan(v):
            return None
        return float(v)
    except Exception:
        return None


def _parse_first_float(v: object) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)

    m = re.search(r"[-+]?\d*\.?\d+", str(v).replace(",", ""))
    if not m:
        return None
    try:
        return float(m.group(0))
    except Exception:
        return None


def _parse_note_date(v: object) -> pd.Timestamp | None:
    if v is None:
        return None
    if isinstance(v, pd.Timestamp):
        return v.normalize()
    if isinstance(v, datetime):
        return pd.Timestamp(v).normalize()
    if isinstance(v, (int, float)):
        # Excel serial date (1900 system)
        if 20000 <= float(v) <= 80000:
            dt = datetime(1899, 12, 30) + timedelta(days=float(v))
            return pd.Timestamp(dt.date())
        return None

    txt = str(v).strip()
    m = re.search(r"(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})", txt)
    if not m:
        return None
    y, mm, dd = map(int, m.groups())
    try:
        return pd.Timestamp(year=y, month=mm, day=dd)
    except Exception:
        return None


def _normalize_note_text(v: object) -> str | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        dt = _parse_note_date(v)
        if dt is not None:
            return dt.strftime("%Y/%m/%d")
        return None
    txt = str(v).strip()
    return txt if txt else None


def _merge_note(note: object, marker: str | None) -> str | None:
    base = _normalize_note_text(note)
    marker_txt = (marker or "").strip()
    if not marker_txt:
        return base
    if base and marker_txt in base:
        return base
    return f"{base} | {marker_txt}" if base else marker_txt


def _load_template_rows(ws) -> list[TemplateRow]:
    rows: list[TemplateRow] = []
    current_item: str | None = None

    for r in range(6, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a not in (None, ""):
            current_item = str(a).strip()

        g = ws.cell(r, 7).value
        if not (isinstance(g, str) and g.startswith("=")):
            continue
        if not current_item:
            continue

        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value
        d = ws.cell(r, 4).value
        e = ws.cell(r, 5).value
        f = ws.cell(r, 6).value
        h = ws.cell(r, 8).value

        rows.append(
            TemplateRow(
                excel_row=r,
                item=current_item,
                unit_price_raw=b,
                unit_price_num=_parse_first_float(b),
                exchange_rate=_to_float(c),
                duty_rate=_to_float(d),
                duty_factor=_to_float(e),
                quantity=_to_float(f) or 0.0,
                note=h,
                note_date=_parse_note_date(h),
            )
        )
    return rows


def _build_alias_map(template_items: list[str], source_products: list[str]) -> dict[str, str | None]:
    src_by_canon: dict[str, list[str]] = {}
    for p in source_products:
        src_by_canon.setdefault(_canon(p), []).append(p)

    manual_canon_alias = {
        _canon("MOCA(China)"): _canon("MOCA"),
        _canon("TT(TMTD)CP"): _canon("TT(CP)"),
        _canon("TT(TMTD)G"): _canon("TT(G)"),
        _canon("CTP(PVI)(G)"): _canon("CTP(G)"),
        _canon("GLYCERIN(무심마스)"): _canon("글리세린(무심마스)"),
    }
    forced_unmapped = {
        _canon("AIBN (FINE POWDER)"),
    }

    aliases: dict[str, str | None] = {}
    for item in template_items:
        c_item = _canon(item)

        if c_item in forced_unmapped:
            aliases[item] = None
            continue

        if c_item in manual_canon_alias and manual_canon_alias[c_item] in src_by_canon:
            aliases[item] = src_by_canon[manual_canon_alias[c_item]][0]
            continue

        # 1) canonical exact match
        if c_item in src_by_canon:
            aliases[item] = src_by_canon[c_item][0]
            continue

        # 2) substring match
        cands: list[str] = []
        for src in source_products:
            c_src = _canon(src)
            if c_item in c_src or c_src in c_item:
                cands.append(src)
        if cands:
            cands.sort(key=lambda x: (abs(len(_canon(x)) - len(c_item)), x))
            aliases[item] = cands[0]
            continue

        # 3) conservative fuzzy fallback
        best_src = None
        best_score = 0.0
        for src in source_products:
            score = SequenceMatcher(None, c_item, _canon(src)).ratio()
            if score > best_score:
                best_score = score
                best_src = src
        aliases[item] = best_src if best_score >= 0.9 else None

    return aliases


def _match_inbound_qty(
    inbound_df: pd.DataFrame,
    product: str,
    row: TemplateRow,
) -> float | None:
    if inbound_df.empty:
        return None

    q = inbound_df[inbound_df["product"] == product].copy()
    if q.empty:
        return None

    if row.note_date is not None:
        q = q[(q["in_date"] >= row.note_date - pd.Timedelta(days=2)) & (q["in_date"] <= row.note_date + pd.Timedelta(days=2))]

    if row.unit_price_num is not None:
        q = q[(q["unit_price"] - row.unit_price_num).abs() <= 1e-3]

    if row.exchange_rate is not None:
        q = q[(q["exchange_rate"] - row.exchange_rate).abs() <= 0.5]

    if q.empty:
        return None
    return float(q["quantity"].sum())


def _num_for_excel(v: float) -> int | float:
    if abs(v - round(v)) < 1e-9:
        return int(round(v))
    return round(v, 3)


def _make_output_path(template_file: Path, output_file: Path) -> tuple[Path, bool]:
    tpl = template_file.resolve(strict=False)
    out = output_file.resolve(strict=False)
    if tpl != out:
        return output_file, False

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    suffix = output_file.suffix or ".xlsx"
    candidate = output_file.with_name(f"{output_file.stem}_자동생성_{ts}{suffix}")
    seq = 1
    while candidate.exists():
        candidate = output_file.with_name(f"{output_file.stem}_자동생성_{ts}_{seq}{suffix}")
        seq += 1
    return candidate, True


def _copy_row_template(ws, src_row: int, dst_row: int) -> None:
    """템플릿 데이터 행의 값/서식을 새 행으로 복사한다."""
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
    for col in range(1, ws.max_column + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)

        if src.has_style:
            dst._style = copy(src._style)
        if src.hyperlink:
            dst._hyperlink = copy(src.hyperlink)
        if src.comment:
            dst.comment = copy(src.comment)

        if isinstance(src.value, str) and src.value.startswith("="):
            dst.value = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
        else:
            dst.value = src.value


def _expand_auto_filter(ws, inserted_after_row: int, amount: int) -> None:
    if not ws.auto_filter.ref:
        return

    min_col, min_row, max_col, max_row = range_boundaries(ws.auto_filter.ref)
    if min_row <= inserted_after_row <= max_row:
        max_row += amount
        ws.auto_filter.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def _insert_template_rows_for_item(ws, item_rows: list[TemplateRow], extra_count: int) -> None:
    """FIFO lot 수가 템플릿 행보다 많을 때 마지막 행 스타일을 복제해 행을 추가한다."""
    if extra_count <= 0:
        return

    first_row = item_rows[0].excel_row
    last_row = item_rows[-1].excel_row
    insert_at = last_row + 1
    merged_ranges = list(ws.merged_cells.ranges)

    for merged in merged_ranges:
        ws.unmerge_cells(str(merged))
    ws.insert_rows(insert_at, amount=extra_count)
    for offset in range(extra_count):
        dst_row = insert_at + offset
        _copy_row_template(ws, last_row, dst_row)
        ws.cell(dst_row, 1).value = None

    for merged in merged_ranges:
        min_col, min_row, max_col, max_row = merged.bounds
        is_item_merge = min_col == max_col == 1 and min_row <= first_row <= max_row
        if is_item_merge:
            continue
        if min_row >= insert_at:
            min_row += extra_count
            max_row += extra_count
        elif max_row >= insert_at:
            max_row += extra_count
        ws.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row,
            end_column=max_col,
        )

    if last_row + extra_count > first_row:
        ws.merge_cells(start_row=first_row, start_column=1, end_row=last_row + extra_count, end_column=1)
    _expand_auto_filter(ws, last_row, extra_count)


def _update_template_header(ws, year: int, start: pd.Timestamp, end: pd.Timestamp) -> None:
    """결과물 템플릿 상단의 기간/작성일 텍스트를 선택 기간에 맞춘다."""
    new_title = f"{year}년 {int(end.month)}월"
    if re.search(r"\d{4}년\s*\d{1,2}월", ws.title) and new_title not in ws.parent.sheetnames:
        ws.title = new_title
    if isinstance(ws["A1"].value, str) and "기초원가" in ws["A1"].value:
        ws["A1"].value = f"{year}년  {int(end.month)}월   기초원가"
    if isinstance(ws["H2"].value, str) and "기간" in ws["H2"].value:
        ws["H2"].value = f"기간: {int(start.month)}월 {int(start.day)}일 ~ {int(end.month)}월 {int(end.day)}일"
    if isinstance(ws["H3"].value, str) and "작성일" in ws["H3"].value:
        write_date = end + pd.Timedelta(days=1)
        ws["H3"].value = f"작성일: {int(write_date.month)}월 {int(write_date.day)}일"


def _template_header_matches_period(ws, year: int, start: pd.Timestamp, end: pd.Timestamp) -> bool:
    title_text = f"{year}년"
    period_text = f"{int(start.month)}월 {int(start.day)}일 ~ {int(end.month)}월 {int(end.day)}일"
    title_ok = title_text in str(ws["A1"].value or "") or title_text in str(ws.title)
    period_ok = period_text in str(ws["H2"].value or "")
    return title_ok and period_ok


def _has_excluded_price_inbound(
    df: pd.DataFrame,
    product: str,
    cutoff: pd.Timestamp,
    *,
    include_korean_supplier: bool = True,
) -> bool:
    """원화/국내 매입으로 제외된 입고가 있는지 확인한다."""
    if df.empty:
        return False

    product_s = df["__product__"].astype(str).str.strip()
    in_qty_source = df["__in_qty__"] if "__in_qty__" in df.columns else pd.Series(0.0, index=df.index)
    in_date_source = df["__in_date__"] if "__in_date__" in df.columns else pd.Series(pd.NaT, index=df.index)
    in_qty = pd.to_numeric(in_qty_source, errors="coerce").fillna(0.0)
    in_date = pd.to_datetime(in_date_source, errors="coerce")
    excluded = pd.Series(False, index=df.index)
    excluded_cols = ["__is_krw_accounting_format__", "__is_high_krw_unit_price__"]
    if include_korean_supplier:
        excluded_cols.append("__is_korean_supplier__")
    for col in excluded_cols:
        if col in df.columns:
            excluded = excluded | df[col].fillna(False).astype(bool)

    mask = (
        (product_s == str(product).strip())
        & (in_qty > 0)
        & in_date.notna()
        & (in_date <= cutoff)
        & excluded
    )
    return bool(mask.any())


def _template_quantity_matches_target(item_rows: list[TemplateRow], target_qty: float) -> bool:
    """템플릿 수량 합계가 판매량과 맞으면 사람이 확정한 lot 배분으로 본다."""
    template_qty = sum(float(tr.quantity or 0.0) for tr in item_rows)
    return abs(template_qty - float(target_qty)) <= max(1e-6, abs(float(target_qty)) * 0.001)


def _template_rows_have_krw_price(item_rows: list[TemplateRow]) -> bool:
    for tr in item_rows:
        raw = "" if tr.unit_price_raw is None else str(tr.unit_price_raw).lower()
        if any(token in raw for token in ("₩", "￦", "krw", "원")):
            return True
        if tr.unit_price_num is not None and tr.unit_price_num >= 500.0:
            return True
    return False


def renew_cost_report_from_template(
    inventory_file: Path,
    template_file: Path,
    output_file: Path,
    year: int,
    start: pd.Timestamp,
    end: pd.Timestamp,
    sheet_name: str | None = None,
) -> dict:
    output_file, output_redirected = _make_output_path(template_file, output_file)

    wb = openpyxl.load_workbook(template_file)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    template_period_matches = _template_header_matches_period(ws, year, start, end)
    _update_template_header(ws, year, start, end)
    template_rows = _load_template_rows(ws)

    if not template_rows:
        raise ValueError("템플릿에서 계산식(G열) 데이터 행을 찾지 못했습니다.")

    df = load_inventory_df(str(inventory_file), year)
    if df.empty:
        raise ValueError(f"재고관리대장에서 {year}년 시트를 찾지 못했습니다.")

    sales_df = (
        df[(df["__event_date__"] >= start) & (df["__event_date__"] <= end)]
        .groupby("__product__", as_index=False)["__sales_qty__"]
        .sum()
        .rename(columns={"__product__": "product", "__sales_qty__": "sales_qty"})
    )
    sales_map = {str(r["product"]): float(r["sales_qty"]) for _, r in sales_df.iterrows()}

    template_items: list[str] = []
    seen = set()
    for tr in template_rows:
        if tr.item not in seen:
            template_items.append(tr.item)
            seen.add(tr.item)

    source_products = sorted(set(df["__product__"].astype(str).str.strip().tolist()))
    alias_map = _build_alias_map(template_items, source_products)

    rows_by_item: dict[str, list[TemplateRow]] = {}
    for tr in template_rows:
        rows_by_item.setdefault(tr.item, []).append(tr)

    unresolved_items: list[str] = []
    overflow_items: list[str] = []
    fallback_items: list[str] = []
    excluded_price_items: list[str] = []
    alloc_cache: dict[str, list[dict]] = {}
    allocs_by_item: dict[str, list[dict]] = {}
    extra_rows_by_item: dict[str, int] = {}
    anchored_items: list[str] = []
    anchored_item_set: set[str] = set()
    changed_rows_set: set[int] = set()
    changed_cells = 0

    for item, item_rows in rows_by_item.items():
        source_product = alias_map.get(item)
        if not source_product:
            unresolved_items.append(item)
            continue

        target_qty = float(sales_map.get(source_product, 0.0))
        has_excluded_price = _has_excluded_price_inbound(df, source_product, end)
        should_anchor = (
            template_period_matches
            or _template_quantity_matches_target(item_rows, target_qty)
        )
        if should_anchor and not _template_rows_have_krw_price(item_rows):
            anchored_items.append(item)
            anchored_item_set.add(item)
            allocs_by_item[item] = []
            continue

        if target_qty <= 1e-9:
            allocs_by_item[item] = []
            continue

        if source_product not in alloc_cache:
            alloc_cache[source_product] = fifo_for_product(
                file_path=inventory_file,
                product=source_product,
                year=year,
                sales_qty=target_qty,
                start=start,
                cutoff=end,
                df=df,
                include_ending_stock_lots=True,
            )
        allocs = alloc_cache[source_product]
        allocs_by_item[item] = allocs

        if not allocs:
            if has_excluded_price:
                excluded_price_items.append(item)
            else:
                fallback_items.append(item)
            continue

        if len(allocs) > len(item_rows):
            overflow_items.append(item)
            extra_rows_by_item[item] = len(allocs) - len(item_rows)

    for item, extra_count in sorted(
        extra_rows_by_item.items(),
        key=lambda x: rows_by_item[x[0]][-1].excel_row,
        reverse=True,
    ):
        _insert_template_rows_for_item(ws, rows_by_item[item], extra_count)

    if extra_rows_by_item:
        template_rows = _load_template_rows(ws)
        rows_by_item = {}
        for tr in template_rows:
            rows_by_item.setdefault(tr.item, []).append(tr)

    assigned_qty: dict[int, object] = {}
    assigned_price: dict[int, object] = {}
    assigned_rate: dict[int, object] = {}
    assigned_note: dict[int, object] = {}
    preserve_formula_rows: set[int] = set()

    for item, item_rows in rows_by_item.items():
        source_product = alias_map.get(item)
        if not source_product:
            for tr in item_rows:
                assigned_qty[tr.excel_row] = _KEEP_CELL
                assigned_price[tr.excel_row] = _KEEP_CELL
                assigned_rate[tr.excel_row] = _KEEP_CELL
                assigned_note[tr.excel_row] = _KEEP_CELL
                preserve_formula_rows.add(tr.excel_row)
            continue

        target_qty = float(sales_map.get(source_product, 0.0))
        allocs = allocs_by_item.get(item, [])

        if item in anchored_item_set:
            for tr in item_rows:
                assigned_qty[tr.excel_row] = tr.quantity
                assigned_price[tr.excel_row] = _KEEP_CELL
                assigned_rate[tr.excel_row] = _KEEP_CELL
                assigned_note[tr.excel_row] = _KEEP_CELL
                preserve_formula_rows.add(tr.excel_row)
            continue

        if target_qty <= 1e-9:
            for tr in item_rows:
                assigned_qty[tr.excel_row] = 0.0
                assigned_price[tr.excel_row] = tr.unit_price_num
                assigned_rate[tr.excel_row] = tr.exchange_rate
                assigned_note[tr.excel_row] = _normalize_note_text(tr.note)
            continue

        if not allocs:
            if item in excluded_price_items:
                marker = "원화단가 제외"
                for tr in item_rows:
                    assigned_qty[tr.excel_row] = 0.0
                    assigned_price[tr.excel_row] = _CLEAR_CELL
                    assigned_rate[tr.excel_row] = _CLEAR_CELL
                    assigned_note[tr.excel_row] = _merge_note(tr.note, marker)
                continue

            # 가격 lot을 못 찾았지만 원화 제외 케이스가 아니면 기존 단가/환율을 유지한다.
            for i, tr in enumerate(item_rows):
                assigned_qty[tr.excel_row] = target_qty if i == 0 else 0.0
                assigned_price[tr.excel_row] = tr.unit_price_num
                assigned_rate[tr.excel_row] = tr.exchange_rate
                assigned_note[tr.excel_row] = _normalize_note_text(tr.note)
            continue

        for i, tr in enumerate(item_rows):
            if i < len(allocs):
                a = allocs[i]
                assigned_qty[tr.excel_row] = "-" if a.get("display_only") else float(a["taken_qty"])
                assigned_price[tr.excel_row] = _to_float(a.get("unit_price"))
                assigned_rate[tr.excel_row] = _to_float(a.get("exchange_rate"))
                if a.get("in_date"):
                    assigned_note[tr.excel_row] = str(a.get("in_date")).replace("-", "/")
                else:
                    assigned_note[tr.excel_row] = _normalize_note_text(tr.note)
            else:
                assigned_qty[tr.excel_row] = 0.0
                assigned_price[tr.excel_row] = tr.unit_price_num
                assigned_rate[tr.excel_row] = tr.exchange_rate
                assigned_note[tr.excel_row] = _normalize_note_text(tr.note)

    for tr in template_rows:
        r = tr.excel_row
        old_b = ws.cell(r, 2).value
        old_c = ws.cell(r, 3).value
        old_f = ws.cell(r, 6).value
        old_h = ws.cell(r, 8).value

        qty = assigned_qty.get(r, tr.quantity)
        price = assigned_price.get(r, tr.unit_price_num)
        rate = assigned_rate.get(r, tr.exchange_rate)
        note = assigned_note.get(r, _normalize_note_text(tr.note))

        if qty is not _KEEP_CELL:
            ws.cell(r, 6).value = qty if isinstance(qty, str) else _num_for_excel(qty)
        if price is _KEEP_CELL:
            pass
        elif price is _CLEAR_CELL:
            ws.cell(r, 2).value = None
        elif price is not None:
            ws.cell(r, 2).value = price
        if rate is _KEEP_CELL:
            pass
        elif rate is _CLEAR_CELL:
            ws.cell(r, 3).value = None
        elif rate is not None:
            ws.cell(r, 3).value = rate
        if note is _KEEP_CELL:
            pass
        elif note is not None:
            ws.cell(r, 8).value = note

        if r not in preserve_formula_rows:
            # G열 수식 표준화 (관세계수 있으면 곱)
            e_val = _to_float(ws.cell(r, 5).value)
            if e_val is not None and e_val > 0:
                ws.cell(r, 7).value = f"=B{r}*C{r}*E{r}"
            else:
                ws.cell(r, 7).value = f"=B{r}*C{r}"

        new_b = ws.cell(r, 2).value
        new_c = ws.cell(r, 3).value
        new_f = ws.cell(r, 6).value
        new_h = ws.cell(r, 8).value

        for old_v, new_v in ((old_b, new_b), (old_c, new_c), (old_f, new_f), (old_h, new_h)):
            if old_v != new_v:
                changed_cells += 1
                changed_rows_set.add(r)

    # 계산식 행 외에 기존 숫자행(예: 합계/메모)은 그대로 둔다.
    # 매핑 실패 품목은 원본값 유지로 처리.
    #
    # 참고: 템플릿 품목이 소스 품목보다 적은 경우 신규 품목 자동 추가는 하지 않는다.
    # (서식 보존 우선)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)

    return {
        "output_file": str(output_file),
        "output_redirected": output_redirected,
        "template_rows": len(template_rows),
        "template_items": len(rows_by_item),
        "mapped_items": len(rows_by_item) - len(unresolved_items),
        "unresolved_items": unresolved_items,
        "fallback_items": fallback_items,
        "excluded_price_items": excluded_price_items,
        "overflow_items": overflow_items,
        "anchored_items": anchored_items,
        "changed_rows": len(changed_rows_set),
        "changed_cells": changed_cells,
    }


def generate_cost_report_without_template(
    inventory_file: Path,
    output_file: Path,
    year: int,
    start: pd.Timestamp,
    end: pd.Timestamp,
    sheet_name: str | None = None,
) -> dict:
    output_file, output_redirected = _make_output_path(inventory_file, output_file)

    df = load_inventory_df(str(inventory_file), year)
    if df.empty:
        raise ValueError(f"재고관리대장에서 {year}년 시트를 찾지 못했습니다.")

    sales_df = (
        df[(df["__event_date__"] >= start) & (df["__event_date__"] <= end)]
        .groupby("__product__", as_index=False)["__sales_qty__"]
        .sum()
        .rename(columns={"__product__": "product", "__sales_qty__": "sales_qty"})
    )
    sales_df["product"] = sales_df["product"].astype(str).str.strip()
    sales_df["sales_qty"] = pd.to_numeric(sales_df["sales_qty"], errors="coerce").fillna(0.0)
    sales_df = sales_df[(sales_df["product"] != "") & (sales_df["sales_qty"] > 0)].copy()
    sales_df = sales_df.sort_values("product", key=lambda s: s.str.lower()).reset_index(drop=True)

    fallback_items: list[str] = []
    overflow_items: list[str] = []
    excluded_price_items: list[str] = []
    rows: list[dict] = []

    for _, sr in sales_df.iterrows():
        product = str(sr["product"])
        target_qty = float(sr["sales_qty"])

        allocs = fifo_for_product(
            file_path=inventory_file,
            product=product,
            year=year,
            sales_qty=target_qty,
            start=start,
            cutoff=end,
            df=df,
            include_ending_stock_lots=True,
        )

        if not allocs:
            if _has_excluded_price_inbound(df, product, end):
                excluded_price_items.append(product)
                continue

            fallback_items.append(product)
            rows.append(
                {
                    "item": product,
                    "unit_price": None,
                    "exchange_rate": None,
                    "duty_rate": 0.0,
                    "duty_factor": 0.0,
                    "quantity": target_qty,
                    "note": "단가 lot 미탐지",
                }
            )
            continue

        for a in allocs:
            rows.append(
                {
                    "item": product,
                    "unit_price": _to_float(a.get("unit_price")),
                    "exchange_rate": _to_float(a.get("exchange_rate")),
                    "duty_rate": 0.0,
                    "duty_factor": 0.0,
                    "quantity": "-" if a.get("display_only") else float(a["taken_qty"]),
                    "note": str(a.get("in_date")).replace("-", "/") if a.get("in_date") else "",
                }
            )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name or "원가결과"

    headers = ["품목", "단가", "환율", "관세율", "관세계수", "수량", "금액", "비고"]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c).value = h

    changed_cells = 0
    for i, row in enumerate(rows, start=2):
        ws.cell(i, 1).value = row["item"]
        changed_cells += 1

        if row["unit_price"] is not None:
            ws.cell(i, 2).value = row["unit_price"]
            changed_cells += 1
        if row["exchange_rate"] is not None:
            ws.cell(i, 3).value = row["exchange_rate"]
            changed_cells += 1

        ws.cell(i, 4).value = row["duty_rate"]
        ws.cell(i, 5).value = row["duty_factor"]
        ws.cell(i, 6).value = row["quantity"] if isinstance(row["quantity"], str) else _num_for_excel(float(row["quantity"]))
        changed_cells += 3

        if row["unit_price"] is not None and row["exchange_rate"] is not None:
            ws.cell(i, 7).value = f"=B{i}*C{i}"
            changed_cells += 1
        if row["note"]:
            ws.cell(i, 8).value = row["note"]
            changed_cells += 1

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)

    return {
        "output_file": str(output_file),
        "output_redirected": output_redirected,
        "template_rows": len(rows),
        "template_items": int(len(sales_df)),
        "mapped_items": int(len(sales_df)),
        "unresolved_items": [],
        "fallback_items": fallback_items,
        "excluded_price_items": excluded_price_items,
        "overflow_items": overflow_items,
        "changed_rows": len(rows),
        "changed_cells": changed_cells,
    }


def _legacy_renew_cost_report_from_template_unused(
    inventory_file: Path,
    template_file: Path,
    output_file: Path,
    year: int,
    start: pd.Timestamp,
    end: pd.Timestamp,
    sheet_name: str | None = None,
) -> dict:
    """
    기존 리스케일 방식 보관용 (사용하지 않음)
    """
    wb = openpyxl.load_workbook(template_file)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    template_rows = _load_template_rows(ws)

    if not template_rows:
        raise ValueError("템플릿에서 계산식(G열) 데이터 행을 찾지 못했습니다.")

    df = load_inventory_df(str(inventory_file), year)
    if df.empty:
        raise ValueError(f"재고관리대장에서 {year}년 시트를 찾지 못했습니다.")

    sales_df = (
        df[(df["__event_date__"] >= start) & (df["__event_date__"] <= end)]
        .groupby("__product__", as_index=False)["__sales_qty__"]
        .sum()
        .rename(columns={"__product__": "product", "__sales_qty__": "sales_qty"})
    )
    sales_map = {str(r["product"]): float(r["sales_qty"]) for _, r in sales_df.iterrows()}

    template_items: list[str] = []
    seen = set()
    for tr in template_rows:
        if tr.item not in seen:
            template_items.append(tr.item)
            seen.add(tr.item)

    source_products = sorted(set(df["__product__"].astype(str).str.strip().tolist()))
    alias_map = _build_alias_map(template_items, source_products)

    rows_by_item: dict[str, list[TemplateRow]] = {}
    for tr in template_rows:
        rows_by_item.setdefault(tr.item, []).append(tr)

    assigned: dict[int, float] = {}
    unresolved_items: list[str] = []

    for item, item_rows in rows_by_item.items():
        source_product = alias_map.get(item)
        if not source_product:
            unresolved_items.append(item)
            for tr in item_rows:
                assigned[tr.excel_row] = tr.quantity
            continue

        target_qty = float(sales_map.get(source_product, 0.0))
        current_total = sum(tr.quantity for tr in item_rows)

        if current_total <= 1e-9 or target_qty <= 1e-9:
            for tr in item_rows:
                assigned[tr.excel_row] = tr.quantity
            continue

        if abs(current_total - target_qty) <= max(1e-9, current_total * 0.01):
            for tr in item_rows:
                assigned[tr.excel_row] = tr.quantity
            continue

        remain = target_qty
        # 기존 분할 비율이 있으면 비율로 리스케일
        for tr in item_rows[:-1]:
            scaled = target_qty * (tr.quantity / current_total)
            assigned[tr.excel_row] = max(0.0, scaled)
            remain -= scaled
        assigned[item_rows[-1].excel_row] = max(0.0, remain)

    for excel_row, qty in assigned.items():
        ws.cell(excel_row, 6).value = _num_for_excel(qty)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)

    return {
        "output_file": str(output_file),
        "template_rows": len(template_rows),
        "template_items": len(rows_by_item),
        "unresolved_items": unresolved_items,
        "carry_over_items": [],
    }
