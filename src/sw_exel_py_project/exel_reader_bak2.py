import re
from pathlib import Path
import pandas as pd
import openpyxl
import unicodedata
from .config.exel_column_map import INVENTORY_COLUMN_MAP


# ======================
# 공통 유틸
# ======================

def _canon(s: object) -> str:
    """공백/괄호/구분기호 제거 + 소문자 정규화"""
    if s is None or (isinstance(s, float) and pd.isna(s)):
        return ""
    s = unicodedata.normalize("NFKC", str(s))
    return re.sub(r"\s+|\(|\)|\[|\]|{|}|/|\\|-|_", "", s).lower()

def _parse_korean_md(s: str, year: int):
    """'MM월 DD일' → Timestamp 변환"""
    if pd.isna(s): 
        return pd.NaT
    if isinstance(s, pd.Timestamp): 
        return s.normalize()
    txt = str(s).strip()
    m = re.match(r"(\d{1,2})\s*월\s*(\d{1,2})\s*일", txt)
    if m:
        try:
            return pd.Timestamp(year=year, month=int(m[1]), day=int(m[2]))
        except: 
            return pd.NaT
    return pd.to_datetime(txt, errors="coerce")

def _norm(s): 
    return "" if pd.isna(s) else str(s).strip().lower()

def _is_unnamed(s): 
    return _norm(s).startswith("unnamed")


# ======================
# 컬럼 선택기
# ======================

def _pick_date_col(df, col_date_lv0):
    for c in df.columns:
        if isinstance(c, tuple) and _norm(c[0]) == _norm(col_date_lv0): 
            return c
        if not isinstance(c, tuple) and _norm(c) == _norm(col_date_lv0): 
            return c
    raise KeyError(f"날짜 컬럼(level=0='{col_date_lv0}') 없음")

def _pick_product_col(df, col_prod_lv0):
    for c in df.columns:
        if isinstance(c, tuple) and _norm(c[0]) == _norm(col_prod_lv0): 
            return c
        if not isinstance(c, tuple) and _norm(c) == _norm(col_prod_lv0): 
            return c
    raise KeyError(f"품명 컬럼(level=0='{col_prod_lv0}') 없음")

def _pick_sales_qty_cols(df: pd.DataFrame):
    """납품/매출 블록의 수량(Kg) 컬럼"""
    chosen = []
    if isinstance(df.columns, pd.MultiIndex):
        for col in df.columns:
            parts = [_canon(x) for x in col]
            if not parts: 
                continue
            top = parts[0]
            if (("납품" not in top) and ("매출" not in top)) or ("입고" in top) or ("수입" in top):
                continue
            rest = parts[1:]
            if any(("수량" in p) or ("kg" in p) for p in rest):
                chosen.append(col)
    else:
        for col in df.columns:
            p = _canon(col)
            if (("납품" in p) or ("매출" in p)) and not (("입고" in p) or ("수입" in p)):
                if ("수량" in p) or ("kg" in p): 
                    chosen.append(col)
    if not chosen:
        raise ValueError("납품/매출 수량 열 없음")
    return chosen

def _pick_inbound_unit_price_cols(df: pd.DataFrame):
    """입고/수입 단가 컬럼"""
    chosen = []
    if isinstance(df.columns, pd.MultiIndex):
        for col in df.columns:
            parts = [_canon(x) for x in col]
            if ("입고" in parts[0] or "수입" in parts[0]) and any("단가" in p for p in parts[1:]):
                chosen.append(col)
    else:
        for col in df.columns:
            p = _canon(col)
            if ("입고" in p or "수입" in p) and "단가" in p: 
                chosen.append(col)
    return chosen

def _pick_inbound_exchange_rate_cols(df: pd.DataFrame):
    """입고/수입 환율 컬럼"""
    chosen = []
    if isinstance(df.columns, pd.MultiIndex):
        for col in df.columns:
            parts = [_canon(x) for x in col]
            if ("입고" in parts[0] or "수입" in parts[0]) and any("환율" in p for p in parts[1:]):
                chosen.append(col)
    else:
        for col in df.columns:
            p = _canon(col)
            if ("입고" in p or "수입" in p) and "환율" in p: 
                chosen.append(col)
    return chosen


# ======================
# 판매 집계
# ======================

def filter_excel_by_ui(result: dict) -> pd.DataFrame:
    file_path = Path(result["file_path"])
    year = int(result["year"])

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheets = [s for s in wb.sheetnames if str(year) in s]
    if not sheets:
        raise ValueError(f"엑셀에 '{year}' 시트 없음")

    df = pd.read_excel(
        file_path, sheet_name=sheets[0],
        engine="openpyxl", header=[2, 3], dtype=str
    )

    date_col = _pick_date_col(df, INVENTORY_COLUMN_MAP["date"])
    prod_col = _pick_product_col(df, INVENTORY_COLUMN_MAP["product_name"])

    df["product"] = df.loc[:, prod_col].astype(str).str.strip()
    df[date_col] = df[date_col].apply(lambda x: _parse_korean_md(x, year))

    start = pd.Timestamp(f"{year}-{int(result['start_month']):02d}-{int(result['start_day']):02d}")
    end   = pd.Timestamp(f"{year}-{int(result['end_month']):02d}-{int(result['end_day']):02d}")
    df = df[(df[date_col] >= start) & (df[date_col] <= end)]

    qty_cols = _pick_sales_qty_cols(df)
    df["__qty__"] = (df[qty_cols]
                     .replace(r"[^\d\.\-]", "", regex=True)
                     .apply(pd.to_numeric, errors="coerce")
                     .fillna(0.0).sum(axis=1))

    out = (df[["product", "__qty__"]]
           .rename(columns={"__qty__": "quantity"}))

    out["product"] = out["product"].replace({"nan": None, "None": None, "": None}).str.strip()
    out = out[out["product"].notna()]

    out = out.groupby("product", as_index=False)["quantity"].sum()
    out = out.sort_values("product", key=lambda s: s.str.lower()).reset_index(drop=True)

    return out


# ======================
# FIFO 관련
# ======================

def load_inbound_history(file_path: Path, year: int, product: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheets = [s for s in wb.sheetnames if str(year) in s]
    if not sheets: 
        return pd.DataFrame()

    df = pd.read_excel(file_path, sheet_name=sheets[0], engine="openpyxl", header=[2, 3], dtype=str)

    date_col = _pick_date_col(df, INVENTORY_COLUMN_MAP["date"])
    prod_col = _pick_product_col(df, INVENTORY_COLUMN_MAP["product_name"])
    df[date_col] = df[date_col].apply(lambda x: _parse_korean_md(x, year))

    # ✅ 입고 수량 (canon 적용)
    qty_cols = []
    for c in df.columns:
        if isinstance(c, tuple):
            parts = [_canon(x) for x in c]
            if ("입고" in parts[0] or "수입" in parts[0]) and any("수량" in p or "kg" in p for p in parts[1:]):
                qty_cols.append(c)

    if not qty_cols: 
        return pd.DataFrame()

    df["__in_qty__"] = (df[qty_cols].replace(r"[^\d\.\-]", "", regex=True)
                                     .apply(pd.to_numeric, errors="coerce")
                                     .fillna(0.0).sum(axis=1))

    # 단가/환율
    price_cols = _pick_inbound_unit_price_cols(df)
    rate_cols  = _pick_inbound_exchange_rate_cols(df)

    unit_price = (df[price_cols].replace(r"[^\d\.\-]","",regex=True)
                                .apply(pd.to_numeric,errors="coerce")
                                .mean(axis=1) if price_cols else None)
    exchange   = (df[rate_cols].replace(r"[^\d\.\-]","",regex=True)
                               .apply(pd.to_numeric,errors="coerce")
                               .mean(axis=1) if rate_cols else None)

    inbound = pd.DataFrame({
        "product": df[prod_col].astype(str).str.strip(),
        "in_date": df[date_col],
        "quantity": df["__in_qty__"],
        "unit_price": unit_price if unit_price is not None else None,
        "exchange_rate": exchange if exchange is not None else None
    })

    return (inbound.query("product == @product and quantity > 0")
                   .dropna(subset=["in_date"])
                   .sort_values("in_date")
                   .reset_index(drop=True))

def fifo_allocate(sales_qty: float, inbound_df: pd.DataFrame):
    allocations, remaining = [], sales_qty
    for _, row in inbound_df.iterrows():
        if remaining <= 0: break
        take = min(remaining, row["quantity"])
        allocations.append({
            "in_date": row["in_date"],
            "taken_qty": take,
            "unit_price": row.get("unit_price"),
            "exchange_rate": row.get("exchange_rate")
        })
        remaining -= take
    return allocations, remaining

def fifo_for_product(file_path: Path, product: str, year: int, sales_qty: float):
    """특정 품목의 판매량을 FIFO 방식으로 입고분과 매칭"""
    # ✅ 해당 연도까지의 모든 입고 불러오기
    inbound_list = []
    for y in range(2000, year + 1):  # 2000년부터 year까지
        inbound = load_inbound_history(file_path, y, product)
        if not inbound.empty:
            inbound_list.append(inbound)

    if not inbound_list:
        raise ValueError(f"{product}: 입고 내역 없음")

    inbound_all = pd.concat(inbound_list, ignore_index=True)
    inbound_all = inbound_all.sort_values("in_date").reset_index(drop=True)

    # ✅ FIFO 배분
    allocations, remaining = fifo_allocate(sales_qty, inbound_all)

    if remaining > 0:
        # 🚨 부족분도 결과에 포함
        allocations.append({
            "in_date": None,
            "taken_qty": remaining,
            "unit_price": None,
            "exchange_rate": None,
            "note": "입고 부족"
        })

    return allocations


# ======================
# 최종 실행
# ======================

def filter_and_fifo(result: dict):
    file_path = Path(result["file_path"])
    year = int(result["year"])

    sales_df = filter_excel_by_ui(result)
    final_results = []
    for _, row in sales_df.iterrows():
        product = row["product"]
        qty = row["quantity"]
        print(f"DEBUG - FIFO 처리중: {product}, 판매수량={qty}")
        allocations = fifo_for_product(file_path, product, year, qty)
        print("DEBUG - FIFO 결과:", allocations)
        final_results.append({
            "product": product,
            "sales_qty": qty,
            "fifo_allocations": allocations
        })

    return final_results

# 조금 잘 되는거 같은데 FIFO가 엉성..