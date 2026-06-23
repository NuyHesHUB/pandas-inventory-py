import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd

from sw_exel_py_project.cost_report import (
    TemplateRow,
    _build_alias_map,
    _match_inbound_qty,
    _parse_note_date,
    generate_cost_report_without_template,
    renew_cost_report_from_template,
)


class CostReportTest(unittest.TestCase):
    def test_parse_note_date_excel_serial(self):
        dt = _parse_note_date(45825)
        self.assertIsNotNone(dt)
        self.assertEqual(str(dt.date()), "2025-06-17")

    def test_parse_note_date_text(self):
        dt = _parse_note_date("2025/07/29 입고(경동)")
        self.assertIsNotNone(dt)
        self.assertEqual(str(dt.date()), "2025-07-29")

    def test_alias_map_canonical_match(self):
        template_items = ["CG-172-25KG", "MOCA(China)"]
        source_products = ["CG 172-25KG", "MOCA"]
        alias = _build_alias_map(template_items, source_products)
        self.assertEqual(alias["CG-172-25KG"], "CG 172-25KG")
        self.assertEqual(alias["MOCA(China)"], "MOCA")

    def test_match_inbound_qty(self):
        inbound_df = pd.DataFrame(
            {
                "product": ["AMBN", "AMBN"],
                "in_date": [pd.Timestamp("2025-07-29"), pd.Timestamp("2025-08-14")],
                "quantity": [16000.0, 16000.0],
                "unit_price": [10.85, 10.85],
                "exchange_rate": [1372.0, 1388.1],
            }
        )
        row = TemplateRow(
            excel_row=12,
            item="AMBN",
            unit_price_raw=10.85,
            unit_price_num=10.85,
            exchange_rate=1372.0,
            duty_rate=0.0,
            duty_factor=0.0,
            quantity=0.0,
            note="2025/07/29",
            note_date=pd.Timestamp("2025-07-29"),
        )
        qty = _match_inbound_qty(inbound_df, "AMBN", row)
        self.assertEqual(qty, 16000.0)

    def test_generate_cost_report_without_template(self):
        df = pd.DataFrame(
            {
                "__event_date__": [pd.Timestamp("2026-01-10")],
                "__product__": ["AMBN"],
                "__sales_qty__": [100.0],
            }
        )
        allocs = [
            {
                "in_date": "2026-01-19",
                "taken_qty": 100.0,
                "unit_price": 10.74,
                "exchange_rate": 1471.5,
            }
        ]

        td_path = Path("artifacts") / "tmp_direct_report_test"
        td_path.mkdir(parents=True, exist_ok=True)
        inventory_file = td_path / "inventory.xlsx"
        output_file = td_path / "out.xlsx"
        inventory_file.write_bytes(b"")

        with patch("sw_exel_py_project.cost_report.load_inventory_df", return_value=df), \
             patch("sw_exel_py_project.cost_report.fifo_for_product", return_value=allocs):
            summary = generate_cost_report_without_template(
                inventory_file=inventory_file,
                output_file=output_file,
                year=2026,
                start=pd.Timestamp("2026-01-01"),
                end=pd.Timestamp("2026-01-31"),
            )

        self.assertTrue(output_file.exists())
        self.assertEqual(summary["template_rows"], 1)
        self.assertEqual(summary["template_items"], 1)
        self.assertEqual(summary["fallback_items"], [])

        wb = openpyxl.load_workbook(output_file)
        ws = wb[wb.sheetnames[0]]
        self.assertEqual(ws.cell(2, 1).value, "AMBN")
        self.assertEqual(float(ws.cell(2, 2).value), 10.74)
        self.assertEqual(float(ws.cell(2, 3).value), 1471.5)
        self.assertEqual(float(ws.cell(2, 6).value), 100.0)

    def test_generate_cost_report_without_template_includes_display_only_lot(self):
        df = pd.DataFrame(
            {
                "__event_date__": [pd.Timestamp("2026-05-10")],
                "__product__": ["AMBN"],
                "__sales_qty__": [16640.0],
            }
        )
        allocs = [
            {
                "in_date": "2026-04-13",
                "taken_qty": 15520.0,
                "unit_price": 10.74,
                "exchange_rate": 1508.32,
            },
            {
                "in_date": "2026-04-27",
                "taken_qty": 1120.0,
                "unit_price": 10.74,
                "exchange_rate": 1479.5,
            },
            {
                "in_date": "2026-05-11",
                "taken_qty": 0.0,
                "unit_price": 10.74,
                "exchange_rate": 1465.4,
                "display_only": True,
            },
        ]

        td_path = Path("artifacts") / "tmp_direct_display_lot_test"
        td_path.mkdir(parents=True, exist_ok=True)
        inventory_file = td_path / "inventory.xlsx"
        output_file = td_path / "out.xlsx"
        inventory_file.write_bytes(b"")

        with patch("sw_exel_py_project.cost_report.load_inventory_df", return_value=df), \
             patch("sw_exel_py_project.cost_report.fifo_for_product", return_value=allocs) as fifo_mock:
            summary = generate_cost_report_without_template(
                inventory_file=inventory_file,
                output_file=output_file,
                year=2026,
                start=pd.Timestamp("2026-04-23"),
                end=pd.Timestamp("2026-05-21"),
            )

        self.assertTrue(fifo_mock.call_args.kwargs["include_ending_stock_lots"])
        self.assertEqual(summary["template_rows"], 3)

        wb = openpyxl.load_workbook(output_file, data_only=False)
        ws = wb[wb.sheetnames[0]]
        self.assertEqual(float(ws.cell(2, 6).value), 15520.0)
        self.assertEqual(float(ws.cell(3, 6).value), 1120.0)
        self.assertEqual(ws.cell(4, 6).value, "-")
        self.assertEqual(float(ws.cell(4, 2).value), 10.74)
        self.assertEqual(float(ws.cell(4, 3).value), 1465.4)
        self.assertEqual(ws.cell(4, 7).value, "=B4*C4")
        self.assertEqual(ws.cell(4, 8).value, "2026/05/11")
        wb.close()

    def test_template_report_expands_rows_with_copied_style(self):
        df = pd.DataFrame(
            {
                "__event_date__": [pd.Timestamp("2026-01-10")],
                "__product__": ["AMBN"],
                "__sales_qty__": [300.0],
            }
        )
        allocs = [
            {
                "in_date": "2026-01-01",
                "taken_qty": 100.0,
                "unit_price": 10.0,
                "exchange_rate": 1400.0,
            },
            {
                "in_date": "2026-01-02",
                "taken_qty": 200.0,
                "unit_price": 11.0,
                "exchange_rate": 1410.0,
            },
        ]

        td_path = Path("artifacts") / "tmp_template_report_test"
        td_path.mkdir(parents=True, exist_ok=True)
        inventory_file = td_path / "inventory.xlsx"
        template_file = td_path / "template.xlsx"
        output_file = td_path / "out.xlsx"
        inventory_file.write_bytes(b"")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2026년 1월"
        ws.merge_cells("A1:H1")
        ws["A1"] = "2025년  8월   기초원가"
        ws["H2"] = "기간: 7월 24일 ~ 8월 25일"
        ws["H3"] = "작성일: 8월 26일"
        headers = ["Item", "수입단가\n(USD/kg)", "수입결제", "관세\n(%)", None, "판매량\n(kg)", "기초원가\n(\\/kg)", "비고"]
        values = ["AMBN", 9.0, 1300.0, 0, 0, 0, "=B6*C6", "old"]
        for col, value in enumerate(headers, start=1):
            ws.cell(5, col).value = value
        for col, value in enumerate(values, start=1):
            ws.cell(6, col).value = value
        ws.row_dimensions[6].height = 40
        ws.cell(6, 1).fill = openpyxl.styles.PatternFill("solid", fgColor="FFFF00")
        ws.auto_filter.ref = "A5:H6"
        wb.save(template_file)
        wb.close()

        with patch("sw_exel_py_project.cost_report.load_inventory_df", return_value=df), \
             patch("sw_exel_py_project.cost_report.fifo_for_product", return_value=allocs):
            summary = renew_cost_report_from_template(
                inventory_file=inventory_file,
                template_file=template_file,
                output_file=output_file,
                year=2026,
                start=pd.Timestamp("2026-01-01"),
                end=pd.Timestamp("2026-01-31"),
            )

        self.assertTrue(output_file.exists())
        self.assertEqual(summary["overflow_items"], ["AMBN"])
        self.assertEqual(summary["template_rows"], 2)

        out_wb = openpyxl.load_workbook(output_file, data_only=False)
        out_ws = out_wb[out_wb.sheetnames[0]]
        self.assertEqual(out_ws.title, "2026년 1월")
        self.assertEqual(out_ws["A1"].value, "2026년  1월   기초원가")
        self.assertEqual(out_ws["H2"].value, "기간: 1월 1일 ~ 1월 31일")
        self.assertEqual(out_ws["H3"].value, "작성일: 2월 1일")
        self.assertEqual(float(out_ws.cell(6, 6).value), 100.0)
        self.assertEqual(float(out_ws.cell(7, 6).value), 200.0)
        self.assertEqual(out_ws.cell(6, 7).value, "=B6*C6")
        self.assertEqual(out_ws.cell(7, 7).value, "=B7*C7")
        self.assertEqual(out_ws.row_dimensions[7].height, 40)
        self.assertIn("A6:A7", [str(r) for r in out_ws.merged_cells.ranges])
        self.assertEqual(out_ws.auto_filter.ref, "A5:H7")
        out_wb.close()

    def test_template_report_writes_display_only_ending_stock_lot(self):
        df = pd.DataFrame(
            {
                "__event_date__": [pd.Timestamp("2026-05-10")],
                "__product__": ["AMBN"],
                "__sales_qty__": [16640.0],
            }
        )
        allocs = [
            {
                "in_date": "2026-04-13",
                "taken_qty": 15520.0,
                "unit_price": 10.74,
                "exchange_rate": 1508.32,
            },
            {
                "in_date": "2026-04-27",
                "taken_qty": 1120.0,
                "unit_price": 10.74,
                "exchange_rate": 1479.5,
            },
            {
                "in_date": "2026-05-11",
                "taken_qty": 0.0,
                "unit_price": 10.74,
                "exchange_rate": 1465.4,
                "display_only": True,
            },
        ]

        td_path = Path("artifacts") / "tmp_template_display_lot_test"
        td_path.mkdir(parents=True, exist_ok=True)
        inventory_file = td_path / "inventory.xlsx"
        template_file = td_path / "template.xlsx"
        output_file = td_path / "out.xlsx"
        inventory_file.write_bytes(b"")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2026년 5월"
        headers = ["Item", "수입단가\n(USD/kg)", "수입결제", "관세\n(%)", None, "판매량\n(kg)", "기초원가\n(\\/kg)", "비고"]
        values = ["AMBN", 10.74, 1508.32, 0, 0, 0, "=B6*C6", ""]
        for col, value in enumerate(headers, start=1):
            ws.cell(5, col).value = value
        for col, value in enumerate(values, start=1):
            ws.cell(6, col).value = value
        ws.auto_filter.ref = "A5:H6"
        wb.save(template_file)
        wb.close()

        with patch("sw_exel_py_project.cost_report.load_inventory_df", return_value=df), \
             patch("sw_exel_py_project.cost_report.fifo_for_product", return_value=allocs) as fifo_mock:
            summary = renew_cost_report_from_template(
                inventory_file=inventory_file,
                template_file=template_file,
                output_file=output_file,
                year=2026,
                start=pd.Timestamp("2026-04-23"),
                end=pd.Timestamp("2026-05-21"),
            )

        self.assertTrue(fifo_mock.call_args.kwargs["include_ending_stock_lots"])
        self.assertEqual(summary["template_rows"], 3)

        out_wb = openpyxl.load_workbook(output_file, data_only=False)
        out_ws = out_wb[out_wb.sheetnames[0]]
        self.assertEqual(float(out_ws.cell(6, 6).value), 15520.0)
        self.assertEqual(float(out_ws.cell(7, 6).value), 1120.0)
        self.assertEqual(out_ws.cell(8, 6).value, "-")
        self.assertEqual(float(out_ws.cell(8, 2).value), 10.74)
        self.assertEqual(float(out_ws.cell(8, 3).value), 1465.4)
        self.assertEqual(out_ws.cell(8, 7).value, "=B8*C8")
        self.assertEqual(out_ws.cell(8, 8).value, "2026/05/11")
        self.assertEqual(out_ws.auto_filter.ref, "A5:H8")
        out_wb.close()

    def test_template_report_preserves_human_anchored_lot(self):
        df = pd.DataFrame(
            {
                "__event_date__": [pd.Timestamp("2025-08-01")],
                "__product__": ["BIIR-232"],
                "__sales_qty__": [60.0],
                "__in_date__": [pd.Timestamp("2025-01-03")],
                "__in_qty__": [60.0],
                "__unit_price__": [15.48],
                "__exchange_rate__": [197.4],
                "__supplier__": ["SIBUR"],
            }
        )
        allocs = [
            {
                "in_date": "2025-01-03",
                "taken_qty": 60.0,
                "unit_price": 15.48,
                "exchange_rate": 197.4,
            }
        ]

        td_path = Path("artifacts") / "tmp_template_anchor_test"
        td_path.mkdir(parents=True, exist_ok=True)
        inventory_file = td_path / "inventory.xlsx"
        template_file = td_path / "template.xlsx"
        output_file = td_path / "out.xlsx"
        inventory_file.write_bytes(b"")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2025년 8월"
        ws.cell(5, 1).value = "Item"
        ws.cell(5, 2).value = "수입단가\n(USD/kg)"
        ws.cell(5, 3).value = "수입결제"
        ws.cell(5, 5).value = "관세계수"
        ws.cell(5, 6).value = "판매량\n(kg)"
        ws.cell(5, 7).value = "기초원가\n(\\/kg)"
        ws.cell(5, 8).value = "비고"
        ws.cell(6, 1).value = "BIIR-232"
        ws.cell(6, 2).value = "¥15.814($2.17)"
        ws.cell(6, 3).value = 198.85
        ws.cell(6, 5).value = 1.05
        ws.cell(6, 6).value = 60
        ws.cell(6, 7).value = "=15.814*C6*E6"
        ws.cell(6, 8).value = "2025/03/12(경동)"
        wb.save(template_file)
        wb.close()

        with patch("sw_exel_py_project.cost_report.load_inventory_df", return_value=df), \
             patch("sw_exel_py_project.cost_report.fifo_for_product", return_value=allocs):
            summary = renew_cost_report_from_template(
                inventory_file=inventory_file,
                template_file=template_file,
                output_file=output_file,
                year=2025,
                start=pd.Timestamp("2025-07-24"),
                end=pd.Timestamp("2025-08-25"),
            )

        self.assertEqual(summary["anchored_items"], ["BIIR-232"])

        out_wb = openpyxl.load_workbook(output_file, data_only=False)
        out_ws = out_wb[out_wb.sheetnames[0]]
        self.assertEqual(out_ws.cell(6, 2).value, "¥15.814($2.17)")
        self.assertEqual(float(out_ws.cell(6, 3).value), 198.85)
        self.assertEqual(float(out_ws.cell(6, 6).value), 60.0)
        self.assertEqual(out_ws.cell(6, 7).value, "=15.814*C6*E6")
        self.assertEqual(out_ws.cell(6, 8).value, "2025/03/12(경동)")
        out_wb.close()

    def test_template_report_clears_krw_excluded_fallback_row(self):
        df = pd.DataFrame(
            {
                "__event_date__": [pd.Timestamp("2025-08-07")],
                "__product__": ["895 END"],
                "__sales_qty__": [18.0],
                "__in_date__": [pd.Timestamp("2025-08-07")],
                "__in_qty__": [18.0],
                "__unit_price__": [37500.0],
                "__exchange_rate__": [None],
                "__supplier__": ["디에이치케미칼"],
                "__is_krw_accounting_format__": [True],
                "__is_high_krw_unit_price__": [True],
                "__is_korean_supplier__": [True],
            }
        )

        td_path = Path("artifacts") / "tmp_template_krw_excluded_test"
        td_path.mkdir(parents=True, exist_ok=True)
        inventory_file = td_path / "inventory.xlsx"
        template_file = td_path / "template.xlsx"
        output_file = td_path / "out.xlsx"
        inventory_file.write_bytes(b"")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "2025년 8월"
        ws.cell(5, 1).value = "Item"
        ws.cell(5, 2).value = "수입단가\n(USD/kg)"
        ws.cell(5, 3).value = "수입결제"
        ws.cell(5, 6).value = "판매량\n(kg)"
        ws.cell(5, 7).value = "기초원가\n(\\/kg)"
        ws.cell(5, 8).value = "비고"
        ws.cell(6, 1).value = "895 END"
        ws.cell(6, 2).value = 37500
        ws.cell(6, 2).number_format = '"₩"#,##0'
        ws.cell(6, 3).value = 1300
        ws.cell(6, 6).value = 18
        ws.cell(6, 7).value = "=B6*C6"
        ws.cell(6, 8).value = "2025/08/07"
        wb.save(template_file)
        wb.close()

        with patch("sw_exel_py_project.cost_report.load_inventory_df", return_value=df), \
             patch("sw_exel_py_project.cost_report.fifo_for_product", return_value=[]):
            summary = renew_cost_report_from_template(
                inventory_file=inventory_file,
                template_file=template_file,
                output_file=output_file,
                year=2025,
                start=pd.Timestamp("2025-08-01"),
                end=pd.Timestamp("2025-08-31"),
            )

        self.assertEqual(summary["excluded_price_items"], ["895 END"])

        out_wb = openpyxl.load_workbook(output_file, data_only=False)
        out_ws = out_wb[out_wb.sheetnames[0]]
        self.assertIsNone(out_ws.cell(6, 2).value)
        self.assertIsNone(out_ws.cell(6, 3).value)
        self.assertEqual(float(out_ws.cell(6, 6).value), 0.0)
        self.assertIn("원화단가 제외", out_ws.cell(6, 8).value)
        out_wb.close()


if __name__ == "__main__":
    unittest.main()
